"""Build the PDFs and the deck.

Run:  py build/build_docs.py      (after build/render_reports.py)
Out:  build/fourth-down-whitepaper.pdf
      build/fourth-down-technical-docs.pdf   (all Diataxis pages in one file)
      build/fourth-down-deck.pdf, docs/fourth-down-deck.pptx, docs/fourth-down-deck.md
      build/validation-report.pdf, build/coach-grades.pdf, build/kicker-grades.pdf

PDFs are printed by headless Microsoft Edge (or Chrome). WeasyPrint needs GTK,
which this machine does not have.
"""
import glob
import html
import json
import os
import re
import shutil
import subprocess
import tempfile

import markdown

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS, OUT = os.path.join(ROOT, "docs"), os.path.join(ROOT, "build")
BROWSERS = [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
            r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
            r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            shutil.which("chromium") or "", shutil.which("google-chrome") or ""]

DOC_CSS = """
@page{size:Letter;margin:18mm 16mm}
body{font:10.5pt/1.5 "Segoe UI",Roboto,Arial,sans-serif;color:#1c1f24}
h1{font-size:21pt;margin:0 0 6pt}h2{font-size:14pt;margin:18pt 0 6pt;border-bottom:1px solid #ddd;padding-bottom:3pt}
h3{font-size:12pt;margin:14pt 0 4pt}h4{font-size:11pt;margin:12pt 0 4pt}
table{border-collapse:collapse;width:100%;margin:8pt 0;font-size:9pt;page-break-inside:auto}
th,td{border:1px solid #ccc;padding:3pt 5pt;text-align:left;vertical-align:top}th{background:#f0efea}
tr{page-break-inside:avoid}
code{font-family:Consolas,monospace;font-size:9pt;background:#f3f2ee;padding:0 2pt}
pre{background:#f3f2ee;padding:6pt;font-size:8.5pt;white-space:pre-wrap}
a{color:#2f6fb5}.pb{page-break-before:always}
"""


def browser():
    for b in BROWSERS:
        if b and os.path.exists(b):
            return b
    raise SystemExit("No Edge or Chrome found to print PDFs.")


def print_pdf(html_text, out_pdf):
    tmp = tempfile.mkdtemp()
    page = os.path.join(tmp, "page.html")
    with open(page, "w", encoding="utf-8") as fh:
        fh.write(html_text)
    subprocess.run([browser(), "--headless=new", "--disable-gpu", "--no-first-run",
                    f"--user-data-dir={os.path.join(tmp, 'profile')}", "--no-pdf-header-footer",
                    f"--print-to-pdf={out_pdf}", "file:///" + page.replace("\\", "/")],
                   check=True, capture_output=True, timeout=120)
    shutil.rmtree(tmp, ignore_errors=True)
    print("wrote", os.path.relpath(out_pdf, ROOT))


def md_html(text, title):
    body = markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])
    return f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title>" \
           f"<style>{DOC_CSS}</style></head><body>{body}</body></html>"


def read(rel):
    return open(os.path.join(ROOT, rel), encoding="utf-8").read()


# ---------------------------------------------------------------- deck
def deck_slides():
    """(title, bullets) for each slide. Figures are read from the pipeline's JSON."""
    v = json.load(open(os.path.join(ROOT, "data", "validation.json")))
    g = json.load(open(os.path.join(ROOT, "data", "coach_grades.json")))
    k = json.load(open(os.path.join(ROOT, "data", "kickers.json")))
    m = json.load(open(os.path.join(ROOT, "data", "model.json")))
    oc = v["coach_agreement"]["option_calibration"]
    lg = g["league_by_season"]
    first = lg[0]
    last = [r for r in lg if r["season"] == max(g["model_fit_seasons"])][0]
    graded = [c for c in g["current"] if c["grade"] != "Incomplete"]
    best = sorted(graded, key=lambda c: c["wp_lost_per_game"])[:3]
    worst = sorted(graded, key=lambda c: -c["wp_lost_per_game"])[:3]
    kc = sorted([x for x in k["kickers"] if x["current"]], key=lambda x: -x["par_per_100"])[:3]
    p = lambda x: f"{100 * x:.0f}%"
    return [
        ("Go for it, kick, or punt?", ["A fourth-down decision engine built on every NFL play since 2014",
                                       f"Version {m['version']}"]),
        ("Fourth down is a choice between three gambles",
         ["Go for it: keep the drive, or hand the other team a short field",
          "Field goal: three likely points, but you give up the ball",
          "Punt: give up the ball, but pin them deep"]),
        ("We ask one question", ["After each choice, what is our chance to win the game?",
                                 "Points alone miss the score and the clock: a tying kick with 4 seconds left is worth far more than 3 points",
                                 "The choice with the best chance to win is the call"]),
        ("What it learned from",
         [f"{m['sample']['plays']:,} plays from {m['seasons'][0]} to {m['seasons'][-1]}",
          "How often teams win from every score, clock and spot on the field",
          "How often teams convert, how often kicks go in, where punts land",
          "How much time each kind of play takes off the clock"]),
        ("Everything that matters goes in",
         ["Always: score, time left, yards to go, field position, both teams' timeouts, the two-minute warning, who gets the ball at halftime",
          "Switch on or off: wind, cold and rain · dome and altitude · how good each team is · your actual kicker"]),
        ("How sure is it?",
         ["We re-built the whole engine 20 times on reshuffled games",
          "If almost all 20 agree, the call is confident",
          "If they split, it is a toss-up, and neither choice is a mistake"]),
        ("Does it work?",
         ["We tested it on two seasons it never saw",
          f"When coaches went for it, it predicted a {p(oc['go']['predicted'])} chance to win; those teams won {p(oc['go']['actual'])}",
          f"Field goals: predicted {p(oc['fg']['predicted'])}, actual {p(oc['fg']['actual'])} · Punts: predicted {p(oc['punt']['predicted'])}, actual {p(oc['punt']['actual'])}"]),
        ("Every coach's call, played twice",
         ["Once the way the coach called it, once the way the engine would",
          "A coach is only charged when the engine is sure he was wrong",
          f"Coaches went for it on {p(first['go_rate'])} of fourth downs in {first['season']} and {p(last['go_rate'])} in {last['season']}",
          f"Closest to the math: {', '.join(c['coach'] for c in best)}",
          f"Furthest from it: {', '.join(c['coach'] for c in worst)}",
          "Trust the order more than the exact numbers: the engine's totals run higher than published research"]),
        ("Kickers matter",
         [f"A replacement-level kicker makes {p(k['replacement']['make_pct_50yd'])} from 50 yards; an average one makes {p(k['league_make_pct_50yd'])}",
          "The engine can use your actual kicker's accuracy and range",
          f"Best current kickers above replacement: {', '.join(x['name'] for x in kc)}"]),
        ("What it can't tell you",
         ["'4th and 1' can mean inches or a full yard; the data can't tell them apart",
          "It can't prove that following it wins more games: no one has run that experiment",
          "Small edges are toss-ups, not mistakes"]),
        ("Read the full story", ["White paper: method, math, results and sources",
                                 "docs/fourth-down-whitepaper.md  ·  build/fourth-down-whitepaper.pdf"]),
    ]


DECK_CSS = """
@page{size:13.333in 7.5in;margin:0}
body{margin:0;font-family:"Segoe UI",Roboto,Arial,sans-serif}
.s{width:13.333in;height:7.5in;box-sizing:border-box;padding:0.9in 1.0in;page-break-after:always;
   background:#f6f5f1;color:#1c1f24;position:relative}
.s h1{font-size:40pt;margin:0 0 0.4in;color:#1f5130}
.s ul{font-size:22pt;line-height:1.45;padding-left:0.4in}
.s li{margin:0 0 0.12in}
.s .bar{position:absolute;left:0;top:0;bottom:0;width:0.25in;background:#2f6b3a}
.s .n{position:absolute;right:0.5in;bottom:0.35in;color:#8a8f98;font-size:12pt}
.t h1{font-size:54pt;margin-top:1.4in}
"""


def build_deck():
    slides = deck_slides()
    md = ["# Fourth Down Engine \u2014 plain-English deck\n",
          "<!-- Generated by build/build_docs.py from data/*.json. Do not edit by hand. -->\n"]
    for i, (title, bullets) in enumerate(slides):
        md.append(f"## {i + 1}. {title}\n")
        md += [f"- {b}" for b in bullets]
        md.append("")
    with open(os.path.join(DOCS, "fourth-down-deck.md"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(md))
    print("wrote docs/fourth-down-deck.md")
    parts = []
    for i, (title, bullets) in enumerate(slides):
        cls = "s t" if i == 0 else "s"
        lis = "".join(f"<li>{html.escape(b)}</li>" for b in bullets)
        parts.append(f"<div class='{cls}'><div class='bar'></div><h1>{html.escape(title)}</h1>"
                     f"<ul>{lis}</ul><div class='n'>{i + 1} / {len(slides)}</div></div>")
    print_pdf(f"<!doctype html><html><head><meta charset='utf-8'><style>{DECK_CSS}</style></head>"
              f"<body>{''.join(parts)}</body></html>", os.path.join(OUT, "fourth-down-deck.pdf"))
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
    except ImportError:
        print("python-pptx not installed; skipped the .pptx")
        return
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for title, bullets in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(0.8), Inches(11.3), Inches(1.2)).text_frame
        tb.word_wrap = True
        tb.text = title
        tb.paragraphs[0].runs[0].font.size = Pt(40)
        tb.paragraphs[0].runs[0].font.bold = True
        tb.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x1F, 0x51, 0x30)
        body = s.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.3), Inches(4.8)).text_frame
        body.word_wrap = True
        for j, b in enumerate(bullets):
            para = body.paragraphs[0] if j == 0 else body.add_paragraph()
            para.text = "• " + b
            para.runs[0].font.size = Pt(22)
            para.space_after = Pt(10)
    prs.save(os.path.join(DOCS, "fourth-down-deck.pptx"))
    print("wrote docs/fourth-down-deck.pptx")


# ---------------------------------------------------------------- docs
def build_tech_docs():
    order = ["docs/README.md", "docs/tutorial/getting-started.md"] + \
        sorted(glob.glob(os.path.join(DOCS, "how-to", "*.md"))) + \
        sorted(glob.glob(os.path.join(DOCS, "reference", "*.md"))) + \
        sorted(glob.glob(os.path.join(DOCS, "explanation", "*.md")))
    chunks = []
    for i, f in enumerate(order):
        rel = os.path.relpath(f, ROOT) if os.path.isabs(f) else f
        chunks.append(("<div class='pb'></div>" if i else "") +
                      markdown.markdown(read(rel), extensions=["tables", "fenced_code", "sane_lists"]))
    doc = f"<!doctype html><html><head><meta charset='utf-8'><style>{DOC_CSS}</style></head><body>{''.join(chunks)}</body></html>"
    print_pdf(doc, os.path.join(OUT, "fourth-down-technical-docs.pdf"))


def main():
    os.makedirs(OUT, exist_ok=True)
    print_pdf(md_html(read("docs/fourth-down-whitepaper.md"), "Fourth Down Engine white paper"),
              os.path.join(OUT, "fourth-down-whitepaper.pdf"))
    build_tech_docs()
    build_deck()
    for name in ("validation-report", "coach-grades", "kicker-grades"):
        print_pdf(md_html(read(f"docs/reports/{name}.md"), name), os.path.join(OUT, f"{name}.pdf"))


if __name__ == "__main__":
    main()
